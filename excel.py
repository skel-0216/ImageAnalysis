import openpyxl
from openpyxl import Workbook

from fish2hist import *

src_path_ = 'C:/Users/kms27/Desktop/GitHub/ImageAnalysis/source/'
experiment_ = 'mackerel#2/'
stage_ = 'stage00/'
size_ = '250/'
position_ = 'stomach/'
brightness_ = 'dark/'
img_num = ['00.png', '01.png', '02.png', '03.png', '04.png', '05.png']


sheet_name = experiment_[:-1]


def write_mackerel(path_, work_sheet):
    work_sheet.append([path_])
    for i in range(1, 6):
        temp_path = path_ + img_num[i]
        temp_list = sum(mackerel2hist_H(cv2.imread(temp_path)).tolist(), [])

        result_list = [img_num[i][:2] + ' mackerel'] + temp_list

        work_sheet.append(result_list)
    work_sheet.append(['', ''])
    return


def write_squid(path_, work_sheet):
    work_sheet.append([path_])
    for i in range(0, 5):
        bgr = ['b', 'g', 'r']
        for j in range(3):
            temp_path = path_ + img_num[i]
            temp_list = sum(squid2hist_BGR(cv2.imread(temp_path))[j].tolist(), [])

            result_list = [img_num[i][:2] + ' squid ' + bgr[j]] + temp_list

            work_sheet.append(result_list)
        work_sheet.append(['', ''])
    work_sheet.append(['', ''])
    work_sheet.append(['', ''])
    work_sheet.append(['', ''])
    return


# 엑셀파일 쓰기
wb = Workbook()

ws1_1 = wb.active

ws1_1.title = 'mackerel#2_dark'

write_mackerel('source/mackerel#2/stage00/250/stomach/dark/', ws1_1)
write_mackerel('source/mackerel#2/stage01/250/stomach/dark/', ws1_1)
write_mackerel('source/mackerel#2/stage00/350/stomach/dark/', ws1_1)
write_mackerel('source/mackerel#2/stage01/350/stomach/dark/', ws1_1)

write_mackerel('source/mackerel#2/stage00/250/back/dark/', ws1_1)
write_mackerel('source/mackerel#2/stage01/250/back/dark/', ws1_1)
write_mackerel('source/mackerel#2/stage00/350/back/dark/', ws1_1)
write_mackerel('source/mackerel#2/stage01/350/back/dark/', ws1_1)

ws1_2 = wb.create_sheet('mackerel#2_bright')

write_mackerel('source/mackerel#2/stage00/250/stomach/bright/', ws1_2)
write_mackerel('source/mackerel#2/stage01/250/stomach/bright/', ws1_2)
write_mackerel('source/mackerel#2/stage00/350/stomach/bright/', ws1_2)
write_mackerel('source/mackerel#2/stage01/350/stomach/bright/', ws1_2)

write_mackerel('source/mackerel#2/stage00/250/back/bright/', ws1_2)
write_mackerel('source/mackerel#2/stage01/250/back/bright/', ws1_2)
write_mackerel('source/mackerel#2/stage00/350/back/bright/', ws1_2)
write_mackerel('source/mackerel#2/stage01/350/back/bright/', ws1_2)


ws2_1 = wb.create_sheet('mackerel#3_dark')

write_mackerel('source/mackerel#3/stage00/250/stomach/dark/', ws2_1)
write_mackerel('source/mackerel#3/stage01/250/stomach/dark/', ws2_1)
write_mackerel('source/mackerel#3/stage00/350/stomach/dark/', ws2_1)
write_mackerel('source/mackerel#3/stage01/350/stomach/dark/', ws2_1)

write_mackerel('source/mackerel#3/stage00/250/back/dark/', ws2_1)
write_mackerel('source/mackerel#3/stage01/250/back/dark/', ws2_1)
write_mackerel('source/mackerel#3/stage00/350/back/dark/', ws2_1)
write_mackerel('source/mackerel#3/stage01/350/back/dark/', ws2_1)

ws2_2 = wb.create_sheet('mackerel#3_bright')

write_mackerel('source/mackerel#3/stage00/250/stomach/bright/', ws2_2)
write_mackerel('source/mackerel#3/stage01/250/stomach/bright/', ws2_2)
write_mackerel('source/mackerel#3/stage00/350/stomach/bright/', ws2_2)
write_mackerel('source/mackerel#3/stage01/350/stomach/bright/', ws2_2)

write_mackerel('source/mackerel#3/stage00/250/back/bright/', ws2_2)
write_mackerel('source/mackerel#3/stage01/250/back/bright/', ws2_2)
write_mackerel('source/mackerel#3/stage00/350/back/bright/', ws2_2)
write_mackerel('source/mackerel#3/stage01/350/back/bright/', ws2_2)


ws3_1 = wb.create_sheet('squid#1_dark')

write_squid('source/squid#1/stage00/head/dark/', ws3_1)
write_squid('source/squid#1/stage01/head/dark/', ws3_1)
write_squid('source/squid#1/stage02/head/dark/', ws3_1)

write_squid('source/squid#1/stage00/body/dark/', ws3_1)
write_squid('source/squid#1/stage01/body/dark/', ws3_1)
write_squid('source/squid#1/stage02/body/dark/', ws3_1)


ws3_2 = wb.create_sheet('squid#1_bright')

write_squid('source/squid#1/stage00/body/bright/', ws3_2)
write_squid('source/squid#1/stage01/body/bright/', ws3_2)
write_squid('source/squid#1/stage02/body/bright/', ws3_2)

write_squid('source/squid#1/stage00/head/bright/', ws3_2)
write_squid('source/squid#1/stage01/head/bright/', ws3_2)
write_squid('source/squid#1/stage02/head/bright/', ws3_2)


wb.save("test.xlsx")
